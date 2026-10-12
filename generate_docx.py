#!/usr/bin/env python3
"""Generate a Word document from readable files and nearby screenshots."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image


SKIP_DIRS = {
    "node_modules", ".git", ".svn", ".hg", "dist", "build", "out",
    ".next", ".nuxt", "__pycache__", ".venv", "venv", "env", "target",
    "bin", "obj", ".idea", ".vscode", ".pytest_cache", "coverage",
    ".parcel-cache", ".turbo", "vendor",
}
SKIP_FILES = {
    "package-lock.json", "yarn.lock", "pnpm-lock.yaml", "composer.lock",
    "Gemfile.lock",
}
IMAGE_EXTS = {"png", "jpg", "jpeg", "gif", "bmp", "webp"}
BINARY_EXTS = {
    "ico", "svg", "tiff", "psd", "mp3", "mp4", "mov", "avi", "wav",
    "flac", "ogg", "webm", "zip", "tar", "gz", "rar", "7z", "bz2",
    "xz", "pdf", "doc", "docx", "xls", "xlsx", "ppt", "pptx", "exe",
    "dll", "so", "dylib", "bin", "class", "jar", "o", "a", "woff",
    "woff2", "ttf", "eot", "otf", "db", "sqlite", "sqlite3", "pyc",
    "lock",
}


def natural_key(value: str) -> list[object]:
    return [(0, int(part)) if part.isdigit() else (1, part.casefold())
            for part in re.split(r"(\d+)", value)]


def included_path(path: Path, root: Path, output: Path) -> bool:
    relative = path.relative_to(root)
    if path.resolve() == output.resolve():
        return False
    if any(part in SKIP_DIRS for part in relative.parts[:-1]):
        return False
    name = path.name
    if name in SKIP_FILES:
        return False
    if name.startswith(".") and name != ".env.example":
        return False
    return True


def read_inputs(root: Path, output: Path):
    code_files = []
    images = []
    for path in sorted(root.rglob("*"), key=lambda p: natural_key(p.as_posix())):
        if not path.is_file() or not included_path(path, root, output):
            continue
        extension = path.suffix.lower().lstrip(".")
        relative = path.relative_to(root)
        folder = relative.parent.as_posix() if relative.parent != Path(".") else "(root)"
        if extension in IMAGE_EXTS:
            try:
                with Image.open(path) as image:
                    image.verify()
                images.append({"path": path, "folder": folder, "filename": path.name})
            except Exception:
                continue
            continue
        if extension in BINARY_EXTS:
            continue
        try:
            content = path.read_text(encoding="utf-8")
        except (UnicodeDecodeError, OSError):
            continue
        sample = content[:8000]
        if "\x00" in sample:
            continue
        controls = sum(ord(char) < 32 and char not in "\t\n\r" for char in sample)
        if sample and controls / len(sample) > 0.01:
            continue
        code_files.append({"path": path, "folder": folder, "filename": path.name,
                           "content": content})
    return code_files, images


def pair_images(code_files, images):
    by_folder_images = defaultdict(list)
    by_folder_code = defaultdict(list)
    for image in images:
        by_folder_images[image["folder"]].append(image)
    for code in code_files:
        by_folder_code[code["folder"]].append(code)

    assigned = defaultdict(list)
    leftovers = []
    base_name = lambda name: Path(name).stem.casefold()
    for folder, folder_images in by_folder_images.items():
        remaining = list(folder_images)
        codes = by_folder_code.get(folder, [])
        for code in codes:
            matches = [image for image in remaining
                       if base_name(image["filename"]) == base_name(code["filename"])]
            if matches:
                assigned[id(code)].extend(matches)
                remaining = [image for image in remaining if image not in matches]
        unmatched_codes = [code for code in codes if not assigned.get(id(code))]
        if len(unmatched_codes) == 1 and remaining:
            assigned[id(unmatched_codes[0])].extend(remaining)
            remaining = []
        leftovers.extend(remaining)
    return assigned, leftovers


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), "Times New Roman")
    r_fonts.set(qn("w:hAnsi"), "Times New Roman")
    r_fonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_paragraph(paragraph, fill="F7F7F7", left_border=True):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    if left_border:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "12")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), "D9D9D9")
        borders.append(left)
        p_pr.append(borders)


def add_code(document, content):
    lines = content.replace("\r\n", "\n").replace("\t", "    ").split("\n")
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        shade_paragraph(paragraph)
        run = paragraph.add_run(line if line else " ")
        set_run_font(run)


def add_output_slot(document, rows=4):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run("Paste output screenshot here")
    set_run_font(run, size=10, italic=True, color=(166, 166, 166))
    for _ in range(max(3, rows - 1)):
        blank = cell.add_paragraph(" ")
        set_run_font(blank.runs[0], size=10)


def add_image(document, image_path: Path):
    with Image.open(image_path) as image:
        image = image.convert("RGB") if image.mode not in ("RGB", "RGBA") else image.copy()
        width, height = image.size
        max_width, max_height = 6.5, 6.0
        scale = min(max_width / (width / 96), max_height / (height / 96), 1)
        image_width = width / 96 * scale
        stream = __import__("io").BytesIO()
        image.save(stream, format="PNG")
        stream.seek(0)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(stream, width=Inches(image_width))
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = caption.add_run(image_path.name)
    set_run_font(run, size=9, italic=True, color=(166, 166, 166))


def generate(root: Path, output: Path, slot_rows: int):
    code_files, images = read_inputs(root, output)
    if not code_files:
        raise SystemExit("No readable code files were found.")
    assigned, leftovers = pair_images(code_files, images)
    by_folder = defaultdict(list)
    for code in code_files:
        by_folder[code["folder"]].append(code)

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal_r_pr = normal._element.get_or_add_rPr()
    normal_fonts = normal_r_pr.rFonts
    if normal_fonts is None:
        normal_fonts = OxmlElement("w:rFonts")
        normal_r_pr.insert(0, normal_fonts)
    normal_fonts.set(qn("w:ascii"), "Times New Roman")
    normal_fonts.set(qn("w:hAnsi"), "Times New Roman")
    normal_fonts.set(qn("w:eastAsia"), "Times New Roman")

    folders = sorted(by_folder, key=natural_key)
    for folder_index, folder in enumerate(folders):
        if folder_index:
            document.add_page_break()
        heading = document.add_paragraph()
        heading.paragraph_format.space_after = Pt(8)
        run = heading.add_run(folder)
        set_run_font(run, size=16, bold=True)
        run.underline = True

        files = sorted(by_folder[folder], key=lambda item: natural_key(item["filename"]))
        for file_index, code in enumerate(files):
            title = document.add_paragraph()
            title.paragraph_format.space_before = Pt(4)
            title.paragraph_format.space_after = Pt(6)
            run = title.add_run(code["filename"])
            set_run_font(run, size=14, bold=True)
            run.underline = True
            add_code(document, code["content"])
            output_title = document.add_paragraph()
            output_title.paragraph_format.space_before = Pt(8)
            output_title.paragraph_format.space_after = Pt(4)
            set_run_font(output_title.add_run("Output:"), bold=True)
            matched = assigned.get(id(code), [])
            if matched:
                for image in matched:
                    add_image(document, image["path"])
            else:
                add_output_slot(document, slot_rows)
            if file_index < len(files) - 1:
                document.add_paragraph()

    if leftovers:
        document.add_page_break()
        heading = document.add_paragraph()
        run = heading.add_run("Unmatched images")
        set_run_font(run, size=16, bold=True)
        run.underline = True
        for image in leftovers:
            label = document.add_paragraph()
            set_run_font(label.add_run(f'{image["folder"]} / {image["filename"]}'), bold=True)
            add_image(document, image["path"])

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Generated {output} from {len(code_files)} readable file(s) and {len(images)} image(s).")


def main(argv: Iterable[str] | None = None):
    parser = argparse.ArgumentParser(description="Create a Word document from code and screenshots.")
    parser.add_argument("--root", type=Path, default=Path.cwd(), help="Folder to collect files from.")
    parser.add_argument("--output", type=Path, default=Path("code-export.docx"), help="Destination .docx file.")
    parser.add_argument("--slot-rows", type=int, default=14, help="Blank rows in screenshot placeholders.")
    args = parser.parse_args(argv)
    root = args.root.resolve()
    output = args.output if args.output.is_absolute() else Path.cwd() / args.output
    if not root.is_dir():
        raise SystemExit(f"Input folder does not exist: {root}")
    generate(root, output.resolve(), args.slot_rows)


if __name__ == "__main__":
    main()
